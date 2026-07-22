#!/usr/bin/env python3
"""
为 5 个企业生成 7.20-7.23 的监测数据 Excel 文件
格式严格按照 upload-excel API 要求：
- 排污口（必需）
- 时间（必需，格式：YYYY-MM-DD HH:mm）
- 只包含企业已申请的污染物
"""

import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import datetime, timedelta
import random
import os

# 企业数据（企业名称 -> 排污口名称 -> 已申请的污染物）
ENTERPRISES = {
    '重庆大学': {
        'outlets': ['cpu-1'],
        'pollutants': ['COD', 'NH3-N', 'TP']
    },
    '重庆医科大学': {
        'outlets': ['cqmu1'],
        'pollutants': ['COD', 'NH3-N', 'TP', 'TN']
    },
    '重庆师范大学': {
        'outlets': ['cqnu-1'],
        'pollutants': ['COD', 'NH3-N', 'TP', 'TN']
    },
    '四川美术学院': {
        'outlets': ['sfai-1'],
        'pollutants': ['COD', 'NH3-N', 'TP', 'TN']
    },
    '重庆科技大学': {
        'outlets': ['cqust-1'],
        'pollutants': ['COD', 'NH3-N', 'TP', 'TN']
    },
}

# 污染物配置（列名 -> 阈值）
POLLUTANT_CONFIGS = {
    'COD': {'threshold': 300, 'min': 10, 'max': 400},
    'NH3-N': {'threshold': 25, 'min': 0.1, 'max': 30},
    'TP': {'threshold': 1.0, 'min': 0.01, 'max': 1.5},
    'TN': {'threshold': 15.0, 'min': 0.1, 'max': 20.0},
}

# 时间范围：7.20-7.23，每天 4 个时间点
DAYS = [20, 21, 22, 23]
HOURS = [6, 12, 18, 0]  # 6:00, 12:00, 18:00, 0:00

def generate_value(pollutant_config):
    """生成随机监测值，80% 正常，15% 预警，5% 超标"""
    rand = random.random()
    threshold = pollutant_config['threshold']
    min_val = pollutant_config['min']
    max_val = pollutant_config['max']
    
    if rand < 0.80:  # 80% 正常（低于阈值的 80%）
        return round(random.uniform(min_val, threshold * 0.8), 3)
    elif rand < 0.95:  # 15% 预警（阈值的 80%-100%）
        return round(random.uniform(threshold * 0.8, threshold), 3)
    else:  # 5% 超标（超过阈值）
        return round(random.uniform(threshold, max_val), 3)

def create_excel_file(enterprise_name, outlet_name, pollutants, filename):
    """创建 Excel 文件"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '监测数据'
    
    # 表头
    headers = ['排污口', '时间'] + pollutants
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    # 数据行
    row = 2
    for day in DAYS:
        for hour in HOURS:
            # 时间格式：YYYY-MM-DD HH:mm
            time_str = f'2026-07-{day:02d} {hour:02d}:00'
            
            # 排污口
            ws.cell(row=row, column=1, value=outlet_name)
            # 时间
            ws.cell(row=row, column=2, value=time_str)
            
            # 污染物数据
            col = 3
            for pollutant_name in pollutants:
                config = POLLUTANT_CONFIGS[pollutant_name]
                value = generate_value(config)
                ws.cell(row=row, column=col, value=value)
                col += 1
            
            row += 1
    
    # 调整列宽
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 18
    for col in range(3, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 12
    
    # 保存文件
    output_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'public')
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    print(f"已生成：{filepath} ({row - 2} 行数据)")

def main():
    """主函数"""
    for enterprise_name, config in ENTERPRISES.items():
        for outlet_name in config['outlets']:
            pollutants = config['pollutants']
            filename = f'{enterprise_name}_monitoring_data.xlsx'
            create_excel_file(enterprise_name, outlet_name, pollutants, filename)
    
    print("\n所有 Excel 文件已生成！")

if __name__ == '__main__':
    main()
